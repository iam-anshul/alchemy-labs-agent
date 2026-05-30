"""LlamaParse-based document parsing utilities.

Returns the full document markdown plus, for every detected table, the rows
and a standalone .xlsx file (as bytes) ready to store in `tables.xlsx_bytes`.
"""
from __future__ import annotations

import asyncio
import io
import uuid
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook

from api.events import EventSink
from config import get_settings


@dataclass
class ParsedTable:
    table_id: str
    page: int
    rows: list[list[Any]]
    n_rows: int
    n_cols: int
    xlsx_bytes: bytes


@dataclass
class ParsedDocument:
    text: str
    pages_text: list[str] = field(default_factory=list)
    tables: list[ParsedTable] = field(default_factory=list)
    n_pages: int = 0


def rows_to_xlsx_bytes(rows: list[list[Any]], sheet_name: str = "table") -> bytes:
    """Write a 2D list of rows to a single-sheet .xlsx and return its bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "table"
    for row in rows:
        ws.append([("" if cell is None else cell) for cell in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_parser():
    from llama_cloud_services import LlamaParse

    settings = get_settings()
    if not settings.llama_parse_key:
        raise RuntimeError("LLAMA_PARSE_KEY is not set in .env")

    return LlamaParse(
        api_key=settings.llama_parse_key,
        parse_mode="parse_page_with_agent",
        high_res_ocr=True,
        adaptive_long_table=True,
        outlined_table_extraction=True,
        verbose=False,
    )


def _collect(result: Any) -> ParsedDocument:
    pages_text: list[str] = []
    tables: list[ParsedTable] = []

    for page in getattr(result, "pages", []):
        page_n = getattr(page, "page", None) or (len(pages_text) + 1)
        page_md = getattr(page, "md", None) or getattr(page, "text", "") or ""
        pages_text.append(page_md)

        for item in getattr(page, "items", []) or []:
            if getattr(item, "type", None) != "table":
                continue
            rows = getattr(item, "rows", None) or []
            if not rows:
                continue
            table_id = f"tbl_{uuid.uuid4().hex[:12]}"
            xlsx = rows_to_xlsx_bytes(rows, sheet_name=f"p{page_n}_{table_id[-6:]}")
            tables.append(
                ParsedTable(
                    table_id=table_id,
                    page=int(page_n),
                    rows=rows,
                    n_rows=len(rows),
                    n_cols=max((len(r) for r in rows), default=0),
                    xlsx_bytes=xlsx,
                )
            )

    full_text = "\n\n".join(pages_text)
    return ParsedDocument(
        text=full_text,
        pages_text=pages_text,
        tables=tables,
        n_pages=len(pages_text),
    )


async def parse_document_async(path: str, sink: EventSink = EventSink()) -> ParsedDocument:
    """Parse a document with LlamaParse (async)."""
    await sink.publish("parse_started", {"path": path})
    parser = _build_parser()
    result = await parser.aparse(path)
    parsed = _collect(result)
    await sink.publish("parse_done", {"n_pages": parsed.n_pages, "n_tables": len(parsed.tables)})
    return parsed


def parse_document(path: str) -> ParsedDocument:
    """Synchronous wrapper around `parse_document_async`."""
    return asyncio.run(parse_document_async(path))
